from pathlib import Path
import re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
root=Path(__file__).resolve().parents[1]; lines=(root/'reports/project_report.md').read_text().splitlines(); doc=Document()
sec=doc.sections[0]; sec.top_margin=Inches(.75); sec.bottom_margin=Inches(.75); sec.left_margin=Inches(.8); sec.right_margin=Inches(.8)
styles=doc.styles; styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(10); styles['Normal'].font.color.rgb=RGBColor.from_string('263746')
for name,color in [('Title','0B2638'),('Heading 1','0B756A'),('Heading 2','0B756A')]: styles[name].font.name='Aptos Display'; styles[name].font.color.rgb=RGBColor.from_string(color)
styles['Title'].font.size=Pt(28); styles['Heading 1'].font.size=Pt(18)
header=sec.header.paragraphs[0]; header.text='SENTINELFLOW  /  NETWORK IDS SIMULATION'; header.style=doc.styles['Caption']
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
run=footer.add_run('Defensive cybersecurity education · Synthetic flow records only   |   ')
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)
in_code=False
for line in lines:
    if line.startswith('```'): in_code=not in_code; continue
    if not line.strip(): continue
    if line.startswith('# '):
        p=doc.add_paragraph(line[2:],style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        p=doc.add_paragraph('Project report · 2026'); p.style=doc.styles['Subtitle']; continue
    if line.startswith('## '): doc.add_heading(line[3:],level=1); continue
    if line.startswith('### '): doc.add_heading(line[4:],level=2); continue
    if in_code:
        p=doc.add_paragraph(); p.style=doc.styles['Caption']; p.paragraph_format.left_indent=Inches(.25); r=p.add_run(line); r.font.name='Consolas'; r.font.size=Pt(9); continue
    if re.match(r'^\d+\. ',line): doc.add_paragraph(re.sub(r'^\d+\. ','',line),style='List Number'); continue
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(6); p.paragraph_format.line_spacing=1.12
    # lightweight bold parser
    parts=re.split(r'(\*\*.*?\*\*|`.*?`)',line)
    for part in parts:
        if part.startswith('**') and part.endswith('**'): p.add_run(part[2:-2]).bold=True
        elif part.startswith('`') and part.endswith('`'):
            r=p.add_run(part[1:-1]); r.font.name='Consolas'; r.font.color.rgb=RGBColor.from_string('0B756A')
        else: p.add_run(part)
# append measured summary table
metrics=__import__('json').loads((root/'reports/evaluation.json').read_text()); doc.add_heading('Measured Evaluation Summary',level=1)
t=doc.add_table(rows=1,cols=2); t.style='Light Shading Accent 1'; t.rows[0].cells[0].text='Metric'; t.rows[0].cells[1].text='Measured value'
for k in ['model','evaluated_records','accuracy','precision','recall','f1','roc_auc','confusion_matrix']:
 c=t.add_row().cells; c[0].text=k.replace('_',' ').title(); c[1].text=str(metrics[k])
doc.save(root/'reports/project_report.docx')
